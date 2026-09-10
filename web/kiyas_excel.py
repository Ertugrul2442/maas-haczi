# -*- coding: utf-8 -*-
"""
EXCEL CIKTISI KIYASI -- openpyxl (masaustu) vs JSZip + duz OOXML (tarayici).

NEDEN ELLE XML: SheetJS'in ucretsiz surumu .xlsx yazarken STIL YAZAMIYOR.
11.09.2026'da olculdu (SheetJS ile yazip openpyxl'e okutarak):

    deger, sayi bicimi, birlestirme, sutun genisligi  -> geciyor
    kalin, dolgu rengi, cerceve, hiza, dondurma       -> GECMIYOR

Bu tablodaki renkler bicim degil BILGI: mavi = o ayin neti SGK kazancindan
hesaplandi, turuncu = kazanc asgarinin disinda ama hesap yine asgariden
yapildi. Memur hangi satirin teyitli hangisinin tahmin oldugunu oradan
goruyor. O yuzden Word'deki yol izlendi: JSZip + duz OOXML.

KIYAS NASIL CALISIYOR (Word kiyasiyla ayni kalip, tek komut, uc asama):
  1. Python ayni senaryolari openpyxl ile uretir.
  2. node web/kiyas_excel.js ayni senaryolari JS ile uretir.
  3. Python IKISINI DE openpyxl ile acar ve hucre hucre kiyaslar:
     deger, tur, sayi bicimi, kalinlik, punto, dolgu rengi, cercevenin dort
     kenari ve rengi, hiza/sarma; ayrica birlesmeler, sutun genislikleri,
     dondurulmus bolme ve sayfa adi.

Kisisel veri ICERMEZ (butun senaryolar sentetik) -> CI'da kosar.

    py -3.13 web/kiyas_excel.py          # kiyasla
    py -3.13 web/kiyas_excel.py --tut    # uretilen dosyalari silme
"""

import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")
KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, KOK)

from openpyxl import load_workbook

import hesap
import disa_aktar

CALISMA = os.path.join(KOK, "web", "kiyas_excel_dosya")
VERI = os.path.join(CALISMA, "senaryolar.json")

ASG25, ASG26 = 26005.50, 33030.00
BOS_BILGI = {"il": "", "daire": "", "dosya_no": "", "borclu": "",
             "isveren_adi": "", "personel": "", "imza_ad": "",
             "imza_unvan": "", "imza_sicil": ""}


def K(isyeri, yil, ay, gun, pek, ad="ORNEK KISI", cikis="", eksik=""):
    return {"kol": "4a", "isyeri": isyeri, "yil": yil, "ay": ay, "gun": gun,
            "giris": "", "cikis": cikis, "eksik_neden": eksik, "ad": ad,
            "pek": pek}


def aylar(isyeri, ilk_yil, n, pek, gun=30, gun_degisken=False):
    out = []
    yil, ay = ilk_yil, 1
    for i in range(n):
        g = (30 - (i % 17)) if gun_degisken else gun
        out.append(K(isyeri, yil, ay, g, pek * g / 30.0))
        ay += 1
        if ay > 12:
            ay, yil = 1, yil + 1
    return out


def senaryolar():
    """Her senaryo excel_yaz()'in AYRI bir dalini zorluyor."""
    d = {}

    # 1) KONTROL GRUBU -- hepsi teyitli (dolgusuz), borc yok.
    d["duz"] = dict(
        kayitlar=aylar("900001", 2026, 3, ASG26), isyeri="900001",
        teblig=[2026, 1, 1], karar=[2026, 4, 28], borc=0.0,
        kusur="cevap-yok", bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"))

    # 2) DOSYA BORCU VAR -> ozet bloguna "Dosya borcu" satiri ekleniyor
    #    (kalin 11 punto, digerleri kalin 12 -- iki AYRI stil).
    d["borclu"] = dict(d["duz"], borc=10000.0)

    # 3) BORC TOPLAMDAN BUYUK -> satir yine var ama borclandirma degisiyor.
    d["borc_yuksek"] = dict(d["duz"], borc=900000.0)

    # 4) MAVI DOLGU -- asgari ustu kazanc, net PEK'ten hesaplaniyor.
    d["mavi"] = dict(
        kayitlar=aylar("900002", 2026, 4, 65000.0), isyeri="900002",
        teblig=[2026, 1, 1], karar=[2026, 5, 28], borc=0.0,
        kusur="cevap-yok", bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"))

    # 5) TURUNCU DOLGU -- tarifesi TANIMLI OLMAYAN yil (2027). net_maas
    #    hata firlatiyor, hesap asgariye dusuyor ama teyit "asgari-ustu"
    #    kaliyor: excel_yaz'in turuncu dali TAM BURADA calisiyor ve baska
    #    hicbir senaryoda calismiyor.
    d["turuncu"] = dict(
        kayitlar=aylar("900003", 2027, 4, ASG26 * 2), isyeri="900003",
        teblig=[2027, 1, 1], karar=[2027, 5, 28], borc=0.0,
        kusur="cevap-yok", bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"))

    # 6) KAZANC BILGISI YOK -> "veri-yok", uyarilarda "!!!" satirlari.
    d["kazanc_yok"] = dict(
        kayitlar=[K("900004", 2026, 1, 30, ASG26),
                  K("900004", 2026, 2, 30, 0.0),
                  K("900004", 2026, 3, 30, ASG26)],
        isyeri="900004", teblig=[2026, 1, 1], karar=[2026, 4, 28], borc=0.0,
        kusur="cevap-yok", bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"))

    # 7) UZUN TUTANAK -- 40 ay. Tutanak hucresi ve ondan sonraki UYARILAR
    #    blogunun satir hesabi kayarsa burada gorunur.
    d["uzun"] = dict(
        kayitlar=aylar("900005", 2022, 40, ASG25), isyeri="900005",
        teblig=[2022, 1, 1], karar=[2025, 5, 28], borc=0.0,
        kusur="cevap-yok", bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"))

    # 8) HER AY FARKLI GUN -> cok uzun tutanak + cok uyari.
    d["degisken_gun"] = dict(
        kayitlar=aylar("900006", 2023, 30, ASG25, gun_degisken=True),
        isyeri="900006", teblig=[2023, 1, 1], karar=[2025, 7, 28], borc=0.0,
        kusur="kesinti-yok", bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"))

    # 9) FORM DOLU -- borclu adi bicimleniyor (son kelime SOYAD), dosya no
    #    ve isveren unvani ust bilgi blogunda.
    d["form_dolu"] = dict(
        d["duz"],
        bilgi=dict(BOS_BILGI, borclu="ahmet ışık yılmaz",
                   isveren_adi="ORNEK GIDA SANAYI LTD STI",
                   dosya_no="2025/10000"))

    # 10) BORCLU ADI BOS -> s.kisi'ye dusuyor (Word'de bu yol bir kez
    #     kor nokta cikmisti).
    d["borclu_bos"] = dict(d["duz"], bilgi=dict(BOS_BILGI))

    # 11) UNLUSUZ UNVAN -- ek_yonelme() burada cokuyordu (10.09.2026).
    #     Tutanak metni Excel'e de giriyor, yani ayni cokme buraya da
    #     yansirdi.
    d["unlusuz_unvan"] = dict(d["duz"], bilgi=dict(BOS_BILGI,
                                                   isveren_adi="MKS"))

    # 12) TEK AY -- en kisa tablo; satir aritmetiginin alt sinir durumu.
    d["tek_ay"] = dict(
        kayitlar=[K("900007", 2026, 2, 30, ASG26)], isyeri="900007",
        teblig=[2026, 2, 1], karar=[2026, 3, 28], borc=0.0,
        kusur="cevap-yok", bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"))

    return d


# --------------------------------------------------------------------------
# Kiyas
# --------------------------------------------------------------------------

class Sayac:
    def __init__(self):
        self.kiyas = 0
        self.fark = 0
        self.ornek = []

    def esit(self, alan, a, b):
        self.kiyas += 1
        if a != b:
            self.fark += 1
            if len(self.ornek) < 25:
                self.ornek.append(f"{alan}:  masaustu={a!r}  web={b!r}")


def _kenar(k):
    """Cercevenin bir kenari: (stil, renk). openpyxl None/degeri karisik
    tutuyor, ikisini de tek bicime indiriyoruz."""
    if k is None:
        return (None, None)
    renk = None
    if k.color is not None:
        renk = k.color.rgb if isinstance(k.color.rgb, str) else str(k.color.rgb)
    return (k.style, renk)


def hucre_dok(c):
    """Bir hucrenin KIYASLANACAK butun ozellikleri."""
    dolgu = None
    if c.fill is not None and c.fill.patternType:
        rgb = c.fill.fgColor.rgb
        dolgu = (c.fill.patternType,
                 rgb if isinstance(rgb, str) else str(rgb))
    return {
        "deger": c.value,
        "tur": c.data_type,
        "bicim": c.number_format,
        "kalin": bool(c.font.bold),
        "punto": float(c.font.size) if c.font.size else None,
        "dolgu": dolgu,
        "cerceve": [_kenar(c.border.left), _kenar(c.border.right),
                    _kenar(c.border.top), _kenar(c.border.bottom)],
        "hiza": (c.alignment.horizontal, c.alignment.vertical,
                 bool(c.alignment.wrap_text)),
    }


def ham_refler(yol):
    """Sayfa XML'indeki BUTUN hucre referanslari (ham, openpyxl'siz).

    NEDEN GEREKLI: openpyxl bos degerli bir hucreyi
    <c r="B5" t="inlineStr"></c> diye YAZIYOR ama geri okurken value=None
    diyor. Asagidaki sayfa_dok o hucreyi (deger yok, stil yok diye) atliyor,
    dolayisiyla "bir taraf bos hucreyi yaziyor, oteki hic yazmiyor" farkini
    GOREMIYOR. 11.09.2026'da tam bu oldu: fark ancak arayuz kiyasinin ham
    XML okuyan kismi sayesinde goruldu. Burasi o kor noktayi kapatiyor.
    """
    with zipfile.ZipFile(yol) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return sorted(re.findall(r'<c r="([A-Z]+\d+)"', xml))


def sayfa_dok(yol):
    ws = load_workbook(yol).active
    hucreler = {}
    for satir in ws.iter_rows():
        for c in satir:
            # MergedCell'in font/fill'i yok; sadece varligi onemli.
            if c.__class__.__name__ == "MergedCell":
                continue
            if c.value is None and c.style_id == 0:
                continue          # gercekten bos hucre
            hucreler[c.coordinate] = hucre_dok(c)
    return {
        "ad": ws.title,
        "hucreler": hucreler,
        "birlesmeler": sorted(str(x) for x in ws.merged_cells.ranges),
        "genislikler": {k: v.width for k, v in ws.column_dimensions.items()},
        "dondurma": ws.freeze_panes,
        "en_son_satir": ws.max_row,
        "en_son_sutun": ws.max_column,
    }


def kiyasla(s, ad, a, b):
    s.esit(f"{ad}.sayfa_adi", a["ad"], b["ad"])
    s.esit(f"{ad}.birlesmeler", a["birlesmeler"], b["birlesmeler"])
    s.esit(f"{ad}.sutun_genislikleri", a["genislikler"], b["genislikler"])
    s.esit(f"{ad}.dondurma", a["dondurma"], b["dondurma"])
    s.esit(f"{ad}.en_son_satir", a["en_son_satir"], b["en_son_satir"])
    s.esit(f"{ad}.en_son_sutun", a["en_son_sutun"], b["en_son_sutun"])

    anahtarlar = sorted(set(a["hucreler"]) | set(b["hucreler"]))
    s.esit(f"{ad}.hucre_sayisi", len(a["hucreler"]), len(b["hucreler"]))
    for k in anahtarlar:
        ha, hb = a["hucreler"].get(k), b["hucreler"].get(k)
        if ha is None or hb is None:
            s.esit(f"{ad}.{k}.var", ha is not None, hb is not None)
            continue
        for alan in ("deger", "tur", "bicim", "kalin", "punto", "dolgu",
                     "cerceve", "hiza"):
            s.esit(f"{ad}.{k}.{alan}", ha[alan], hb[alan])


def main():
    tut = "--tut" in sys.argv

    if os.path.isdir(CALISMA):
        shutil.rmtree(CALISMA)
    os.makedirs(os.path.join(CALISMA, "py"))
    os.makedirs(os.path.join(CALISMA, "js"))

    ayarlar = hesap.Ayarlar.yukle(os.path.join(KOK, "ayarlar.json"))
    sen = senaryolar()

    # ---- Asama 1: openpyxl tarafi
    dis_veri = {"senaryolar": []}
    for ad, g in sen.items():
        kayitlar = [hesap.Kayit(**k) for k in g["kayitlar"]]
        s = hesap.hesapla(kayitlar, g["isyeri"],
                          datetime.date(*g["teblig"]),
                          datetime.date(*g["karar"]),
                          g["borc"], ayarlar, "ORNEK KISI", g["kusur"])
        disa_aktar.excel_yaz(os.path.join(CALISMA, "py", ad + ".xlsx"),
                             s, g["bilgi"])
        dis_veri["senaryolar"].append({
            "ad": ad, "kayitlar": g["kayitlar"], "isyeri": g["isyeri"],
            "teblig": {"yil": g["teblig"][0], "ay": g["teblig"][1],
                       "gun": g["teblig"][2]},
            "karar": {"yil": g["karar"][0], "ay": g["karar"][1],
                      "gun": g["karar"][2]},
            "borc": g["borc"], "kisi": "ORNEK KISI", "kusur": g["kusur"],
            "bilgi": g["bilgi"],
        })
    with open(VERI, "w", encoding="utf-8") as f:
        json.dump(dis_veri, f, ensure_ascii=False)

    # ---- Asama 2: JS tarafi
    p = subprocess.run(["node", os.path.join(KOK, "web", "kiyas_excel.js"), VERI],
                       capture_output=True, text=True, encoding="utf-8", cwd=KOK)
    if p.returncode != 0:
        print(p.stdout)
        print(p.stderr, file=sys.stderr)
        print(">>> JS tarafi .xlsx uretemedi.")
        sys.exit(1)

    # ---- Asama 3: ikisini de openpyxl ile acip kiyasla
    s = Sayac()
    for ad in sen:
        py_yol = os.path.join(CALISMA, "py", ad + ".xlsx")
        js_yol = os.path.join(CALISMA, "js", ad + ".xlsx")
        # Once ham hucre listesi: bos hucrelerin varligi ancak boyle gorulur.
        s.esit(f"{ad}.ham_hucre_refleri", ham_refler(py_yol), ham_refler(js_yol))
        kiyasla(s, ad, sayfa_dok(py_yol), sayfa_dok(js_yol))

    print("=" * 72)
    print("EXCEL CIKTISI KIYASI  --  openpyxl (masaustu)  vs  JSZip (tarayici)")
    print("=" * 72)
    print(f"  senaryo : {len(sen)}")
    print(f"  kiyas   : {s.kiyas}")
    print(f"  fark    : {s.fark}")
    if s.ornek:
        print("\nIlk farklar:")
        for o in s.ornek:
            print("  - " + o)
    print()
    print(">>> IKI YOL AYNI TABLOYU URETIYOR." if s.fark == 0
          else ">>> AYRISMA VAR -- web surumunun Excel ciktisi masaustunden FARKLI.")

    if not tut:
        shutil.rmtree(CALISMA)
    else:
        print(f"\nDosyalar durdu: {CALISMA}")
    sys.exit(0 if s.fark == 0 else 1)


if __name__ == "__main__":
    main()
