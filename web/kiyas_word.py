# -*- coding: utf-8 -*-
"""
AYNI tutanagi iki yoldan uretir ve YAPISAL olarak kiyaslar:

    masaustu : disa_aktar.word_yaz()  -> python-docx
    tarayici : web/disa_aktar.js      -> duz OOXML + JSZip

Sonra IKISINI DE python-docx ile acip paragraf paragraf, hucre hucre
karsilastirir: metin, hiza, punto, kalinlik, paragraf araligi, girinti,
sayfa kirilmasi, sayfa olculeri, kenar bosluklari, tablo sutun genislikleri.

NEDEN BU KADAR AYRINTILI: bu projede "sayi dogru, metin yanlis" turu sessiz
hata uc kez yasandi (CLAUDE.md). Ustelik burada ek bir risk var -- PUNTO
KADEMELERI Word'e saydirilarak olculdu ve python-docx'in bos sablonunun
sayfa olcusune (LETTER, 1,15 satir araligi) bagli. JS iskeleti A4 ya da tek
satir araligi uretirse tutanak SESSIZCE 2. sayfaya tasar ve bunu kimse
gormez. O yuzden sayfa olcusu ve satir araligi da kiyaslaniyor.

Senaryolarin hepsi UYDURMA -- kisisel veri icermez, CI'da kosar.

Kullanim:
    py -3.13 web/kiyas_word.py            (uc asamayi da kendi calistirir)
    py -3.13 web/kiyas_word.py --tut      (uretilen .docx dosyalarini silme)
"""

import json
import os
import shutil
import subprocess
import sys

sys.stdout.reconfigure(encoding="utf-8")
KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, KOK)

import hesap
import disa_aktar

CALISMA = os.path.join(KOK, "web", "kiyas_word_dosya")
VERI = os.path.join(CALISMA, "senaryolar.json")


# --------------------------------------------------------------------------
# Senaryolar -- her biri baska bir kod yolunu zorluyor
# --------------------------------------------------------------------------

def kayit(yil, ay, gun=30, pek=33030.00, isyeri="900001", ad="ORNEK KISI"):
    return {"kol": "4a", "isyeri": isyeri, "yil": yil, "ay": ay, "gun": gun,
            "giris": "", "cikis": "", "eksik_neden": "", "ad": ad, "pek": pek}


def aylar(baslangic_yil, adet, gun=30, pek=33030.00, gun_degisken=False):
    cikan, y, a = [], baslangic_yil, 1
    for i in range(adet):
        g = (i % 28) + 1 if gun_degisken else gun
        cikan.append(kayit(y, a, g, pek))
        a += 1
        if a > 12:
            a, y = 1, y + 1
    return cikan


BOS_BILGI = {"il": "ornekil", "daire": "3", "dosya_no": "2026/1"}


def senaryolar():
    d = {}

    # 1) EN KISA -- kontrol grubu. Punto en ust kademede (10,5) kalmali.
    d["kisa"] = dict(kayitlar=aylar(2026, 2), teblig=[2026, 1, 1],
                     karar=[2026, 3, 28], borc=0.0, kusur="cevap-yok",
                     bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"), ek=False)

    # 2) EK TABLOLU -- ikinci sayfada hesap dokumu, sayfa kirilmasi,
    #    sabit sutun genislikleri, TableGrid kenarliklari.
    d["kisa_ek"] = dict(d["kisa"], ek=True)

    # 3) PUNTO DUSUYOR -- 40 ay. Kademe 10,5'ten asagi inmeli.
    d["uzun"] = dict(kayitlar=aylar(2022, 40), teblig=[2022, 1, 1],
                     karar=[2025, 5, 28], borc=0.0, kusur="cevap-yok",
                     bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"), ek=False)

    # 4) TEK SAYFAYA SIGMIYOR -- her ayin gunu farkli oldugu icin tutanak
    #    cumlesi patliyor. sigar=False donmeli, iki taraf da AYNI demeli.
    d["sigmayan"] = dict(kayitlar=aylar(2021, 60, gun_degisken=True),
                         teblig=[2021, 1, 1], karar=[2026, 1, 28], borc=0.0,
                         kusur="cevap-yok",
                         bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"),
                         ek=False)

    # 5) DOSYA BORCU GIRILDI -- "dosya borcunun ... olduğu görülmekle" cumlesi
    #    ve ek tablodaki "Dosya borcu" satiri devreye giriyor.
    d["borclu_ek"] = dict(kayitlar=aylar(2026, 3), teblig=[2026, 1, 1],
                          karar=[2026, 4, 28], borc=10000.0,
                          kusur="cevap-yok",
                          bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"),
                          ek=True)

    # 6) KUSUR = KESINTI-YOK -- giris paragrafinin cumlesi degisiyor.
    #    Bu cumle IKI AYRI YERDE uretiliyor (hesap.tutanak_metni ve
    #    disa_aktar); birini degistirip otekini unutmak yasanmis bir hata.
    d["kesinti_yok"] = dict(d["kisa"], kusur="kesinti-yok")

    # 7) UNLUSUZ UNVAN -- ek_yonelme() burada 10.09.2026'da COKUYORDU.
    d["unlusuz_unvan"] = dict(d["kisa"],
                              bilgi=dict(BOS_BILGI, isveren_adi="MKS"))

    # 8) XML KACISLARI -- python-docx bunlari lxml ile kaciriyor, JS elle.
    #    Kacirilmazsa uretilen .docx BOZUK olur ve Word acmaz.
    d["ozel_karakter"] = dict(d["kisa"],
                              bilgi=dict(BOS_BILGI,
                                         isveren_adi="A & B <Ltd> \"Şti\"",
                                         borclu="o'neill & sons",
                                         personel="x <y> & z",
                                         imza_ad="a & b",
                                         imza_unvan="İcra Müdürü <vekil>",
                                         imza_sicil="1<2>3"))

    # 9) IMZA BLOGU DOLU -- personel + uc imza satiri.
    d["imzali"] = dict(d["kisa"],
                       bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD",
                                  personel="ayşe demir",
                                  imza_ad="mehmet ışık yılmaz",
                                  imza_unvan="İcra Müdür Yardımcısı",
                                  imza_sicil="123456"))

    # 10) ISVEREN ADI BOS -- "900001 sicil numarali isyeri" yedegi.
    d["isveren_adi_bos"] = dict(d["kisa"], bilgi=dict(BOS_BILGI))

    # 11) KAZANC BILGISI YOK -- net_kaynak "veri-yok", ek tabloda "!!!"
    #     satiri ve aciklama metni.
    d["veri_yok_ek"] = dict(
        kayitlar=[kayit(2026, 1), kayit(2026, 2, pek=0.0), kayit(2026, 3)],
        teblig=[2026, 1, 1], karar=[2026, 4, 28], borc=0.0,
        kusur="cevap-yok", bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"),
        ek=True)

    # 12) ASGARI USTU -- net_kaynak "pek", her ay farkli net ve 1/4.
    d["asgari_ustu_ek"] = dict(
        kayitlar=aylar(2026, 6, pek=65000.00), teblig=[2026, 1, 1],
        karar=[2026, 7, 28], borc=0.0, kusur="cevap-yok",
        bilgi=dict(BOS_BILGI, isveren_adi="ORNEK LTD"), ek=True)

    # 13) UZUN + EK -- tablo birden fazla sayfaya yayiliyor.
    d["uzun_ek"] = dict(d["uzun"], ek=True)

    # 14) BASLIK ALANLARI BOS -- il/daire/dosya_no yoksa o paragraflar
    #     HIC yazilmamali (bos paragraf birakmak sayfayi kaydirir).
    d["baslik_bos"] = dict(d["kisa"], bilgi={"isveren_adi": "ORNEK LTD"})

    # 15) DAIRE YAZIYLA -- daire_bicimle() rakam degil metin yolu.
    d["daire_yaziyla"] = dict(d["kisa"],
                              bilgi=dict(BOS_BILGI,
                                         daire="Konya 3. İcra Dairesi",
                                         il="ısparta",
                                         isveren_adi="ORNEK LTD"))

    return d


# --------------------------------------------------------------------------
# Belgeyi yapisal olarak dokme -- iki tarafta da AYNI fonksiyon kullaniliyor
# --------------------------------------------------------------------------

def _hiza(p):
    a = p.alignment
    return "yok" if a is None else str(a)


def _uzunluk(x):
    return "yok" if x is None else str(x)


def _kalin(r):
    """
    DIKKAT: bool(r.bold) YAZMA. python-docx'te r.bold uc degerli --
    True (<w:b/>), False (<w:b w:val="0"/>) ve None (hic yok). bool() ile
    sikistirilinca False ile None ayni gorunuyor ve "kalin olmayan run'da
    <w:b w:val=0> yazilmiyor" mutasyonu HIC yakalanmiyordu (10.09.2026).
    """
    return "yok" if r.bold is None else str(r.bold)


def _bosluk_koru(r):
    """
    <w:t xml:space="preserve"> var mi. lxml bunu okurken metne yansitmiyor
    (bosluk zaten .text icinde duruyor), ama WORD yansitiyor: oznitelik
    yoksa "Karar : " sondaki boslugu kirpilir ve cumleler bitisir. Yani
    r.text kiyasi bu hatayi GORMUYOR, ozniteligin kendisine bakmak sart.
    """
    from docx.oxml.ns import qn
    for t in r._element.findall(qn("w:t")):
        if t.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve":
            return True
    return False


# Sablonun sayfa davranisini belirleyen ama python-docx API'siyle
# OKUNAMAYAN degerler. En onemlisi docDefaults'taki satir araligi: 1,15
# (line=276) yerine tek (240) olursa sayfaya daha cok metin sigar ve
# PUNTO_KADEME esikleri gecersiz kalir -- tutanak sessizce 2. sayfaya tasar.
# 10.09.2026'da olculdu: bu mutasyon eskiden HIC yakalanmiyordu.
_ILGILI_STIL = [
    ("varsayilan_punto", r'<w:rPrDefault>.*?<w:sz w:val="(\d+)"'),
    ("varsayilan_bosluk", r'<w:pPrDefault>.*?<w:spacing w:after="(\d+)"'),
    ("varsayilan_satir_araligi", r'<w:pPrDefault>.*?<w:spacing[^/]*w:line="(\d+)"'),
    ("varsayilan_satir_kurali", r'<w:pPrDefault>.*?w:lineRule="(\w+)"'),
]


def stil_dok(yol):
    import re
    import zipfile
    x = zipfile.ZipFile(yol).read("word/styles.xml").decode("utf-8")
    d = {}
    for ad, kalip in _ILGILI_STIL:
        m = re.search(kalip, x, re.S)
        d[ad] = m.group(1) if m else "yok"

    n = re.search(r'w:styleId="Normal".*?</w:style>', x, re.S)
    ic = n.group(0) if n else ""
    m = re.search(r'w:ascii="([^"]*)"', ic)
    d["normal_yazitipi"] = m.group(1) if m else "yok"
    m = re.search(r'<w:sz w:val="(\d+)"', ic)
    d["normal_punto"] = m.group(1) if m else "yok"

    t = re.search(r'w:styleId="TableGrid".*?</w:style>', x, re.S)
    tic = t.group(0) if t else ""
    d["tablo_stili_var"] = bool(tic)
    for kenar in ("top", "left", "bottom", "right", "insideH", "insideV"):
        m = re.search(r'<w:tblBorders>.*?<w:%s w:val="(\w+)" w:sz="(\d+)"' % kenar,
                      tic, re.S)
        d["tablo_kenar_" + kenar] = (m.group(1) + "/" + m.group(2)) if m else "yok"
    m = re.search(r'<w:tblCellMar>.*?<w:left w:w="(\d+)"', tic, re.S)
    d["tablo_hucre_bosluk_sol"] = m.group(1) if m else "yok"
    return d


def belge_dok(yol):
    from docx import Document
    d = Document(yol)

    st = d.styles["Normal"]
    bolum = d.sections[0]

    paragraflar = []
    for p in d.paragraphs:
        pf = p.paragraph_format
        paragraflar.append({
            "metin": p.text,
            "hiza": _hiza(p),
            "bosluk_sonra": _uzunluk(pf.space_after),
            "bosluk_once": _uzunluk(pf.space_before),
            "girinti": _uzunluk(pf.first_line_indent),
            "sayfa_kirilmasi": bool(pf.page_break_before),
            "kosular": [{"metin": r.text, "kalin": _kalin(r),
                         "bosluk_koru": _bosluk_koru(r),
                         "punto": _uzunluk(r.font.size)} for r in p.runs],
        })

    tablolar = []
    for t in d.tables:
        satirlar = []
        for tr in t.rows:
            hucreler = []
            for tc in tr.cells:
                kosular = []
                for tp in tc.paragraphs:
                    for r in tp.runs:
                        kosular.append({"metin": r.text, "kalin": _kalin(r),
                                        "bosluk_koru": _bosluk_koru(r),
                                        "punto": _uzunluk(r.font.size)})
                hucreler.append({"metin": tc.text, "kosular": kosular})
            satirlar.append(hucreler)
        tablolar.append({
            "satir": len(t.rows), "sutun": len(t.columns),
            "sutun_genislik": [_uzunluk(c.width) for c in t.columns],
            "satirlar": satirlar,
        })

    return {
        "sablon": stil_dok(yol),
        "stil_yazitipi": st.font.name,
        "stil_punto": _uzunluk(st.font.size),
        "sayfa_en": _uzunluk(bolum.page_width),
        "sayfa_boy": _uzunluk(bolum.page_height),
        "kenar_ust": _uzunluk(bolum.top_margin),
        "kenar_alt": _uzunluk(bolum.bottom_margin),
        "kenar_sol": _uzunluk(bolum.left_margin),
        "kenar_sag": _uzunluk(bolum.right_margin),
        "paragraflar": paragraflar,
        "tablolar": tablolar,
    }


# --------------------------------------------------------------------------
# Kiyas
# --------------------------------------------------------------------------

class Sayac:
    def __init__(self):
        self.kiyas = 0
        self.fark = 0
        self.ornek = []

    def esit(self, yol, a, b):
        self.kiyas += 1
        if a != b:
            self.fark += 1
            if len(self.ornek) < 25:
                self.ornek.append(f"{yol}\n      python-docx: {a!r}\n      js         : {b!r}")


def kiyasla(s, etiket, a, b):
    if isinstance(a, dict) and isinstance(b, dict):
        for anahtar in sorted(set(a) | set(b)):
            kiyasla(s, f"{etiket}.{anahtar}", a.get(anahtar, "(yok)"),
                    b.get(anahtar, "(yok)"))
    elif isinstance(a, list) and isinstance(b, list):
        s.esit(f"{etiket}.uzunluk", len(a), len(b))
        for i in range(min(len(a), len(b))):
            kiyasla(s, f"{etiket}[{i}]", a[i], b[i])
    else:
        s.esit(etiket, a, b)


def main():
    tut = "--tut" in sys.argv

    if os.path.isdir(CALISMA):
        shutil.rmtree(CALISMA)
    os.makedirs(os.path.join(CALISMA, "py"))
    os.makedirs(os.path.join(CALISMA, "js"))

    ayarlar = hesap.Ayarlar.yukle(os.path.join(KOK, "ayarlar.json"))
    sen = senaryolar()

    # ---- Asama 1: python-docx tarafi
    import datetime
    dis_veri = {"senaryolar": []}
    py_sonuc = {}
    for ad, g in sen.items():
        kayitlar = [hesap.Kayit(**k) for k in g["kayitlar"]]
        s = hesap.hesapla(
            kayitlar, g["kayitlar"][0]["isyeri"],
            datetime.date(*g["teblig"]), datetime.date(*g["karar"]),
            g["borc"], ayarlar, "ORNEK KISI", g["kusur"])
        yol = os.path.join(CALISMA, "py", ad + ".docx")
        _, punto, sigar = disa_aktar.word_yaz(yol, s, g["bilgi"], ek=g["ek"])
        py_sonuc[ad] = {"punto": punto, "sigar": sigar}
        dis_veri["senaryolar"].append({
            "ad": ad, "kayitlar": g["kayitlar"],
            "isyeri": g["kayitlar"][0]["isyeri"],
            "teblig": {"yil": g["teblig"][0], "ay": g["teblig"][1],
                       "gun": g["teblig"][2]},
            "karar": {"yil": g["karar"][0], "ay": g["karar"][1],
                      "gun": g["karar"][2]},
            "borc": g["borc"], "kisi": "ORNEK KISI", "kusur": g["kusur"],
            "bilgi": g["bilgi"], "ek": g["ek"],
        })
    with open(VERI, "w", encoding="utf-8") as f:
        json.dump(dis_veri, f, ensure_ascii=False)

    # ---- Asama 2: JS tarafi
    p = subprocess.run(["node", os.path.join(KOK, "web", "kiyas_word.js"), VERI],
                       capture_output=True, text=True, encoding="utf-8", cwd=KOK)
    if p.returncode != 0:
        print(p.stdout)
        print(p.stderr, file=sys.stderr)
        print(">>> JS tarafi .docx uretemedi.")
        sys.exit(1)
    js_sonuc = json.loads(p.stdout)

    # ---- Asama 3: ikisini de python-docx ile acip kiyasla
    s = Sayac()
    for ad in sen:
        s.esit(f"{ad}.punto", py_sonuc[ad]["punto"], js_sonuc[ad]["punto"])
        s.esit(f"{ad}.tek_sayfaya_sigar", py_sonuc[ad]["sigar"], js_sonuc[ad]["sigar"])
        a = belge_dok(os.path.join(CALISMA, "py", ad + ".docx"))
        b = belge_dok(os.path.join(CALISMA, "js", ad + ".docx"))
        kiyasla(s, ad, a, b)

    print("=" * 72)
    print("WORD CIKTISI KIYASI  --  python-docx (masaustu)  vs  JSZip (tarayici)")
    print("=" * 72)
    print(f"  senaryo : {len(sen)}")
    print(f"  kiyas   : {s.kiyas}")
    print(f"  fark    : {s.fark}")
    if s.ornek:
        print("\nIlk farklar:")
        for o in s.ornek:
            print("  - " + o)
    print()
    print(">>> IKI YOL AYNI BELGEYI URETIYOR." if s.fark == 0
          else ">>> AYRISMA VAR -- web surumunun Word ciktisi masaustunden FARKLI.")

    if not tut:
        shutil.rmtree(CALISMA)
    else:
        print(f"\nDosyalar durdu: {CALISMA}")
    sys.exit(0 if s.fark == 0 else 1)


if __name__ == "__main__":
    main()
